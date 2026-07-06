"""
╔══════════════════════════════════════════════════════════════════╗
║   Crossfire Legends - Facebook Page Scraper & Analyzer  v1.0     ║
║   Purpose   : Fetch & process posts and comments from Facebook   ║
║               official fanpage.                                  ║
║   Features  : - Robust UTF-8 encoding configuration              ║
║               - Unified Sentiment & Error Categorization         ║
║               - Export results to structured Excel reports       ║
╚══════════════════════════════════════════════════════════════════╝
"""

import os
import re
import sys
import time
import logging
from datetime import datetime
import pandas as pd
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

# ─────────────────────────────────────────────────────────────────
# 0. PREVENT MOJIBAKE & SETUP LOGGING
# ─────────────────────────────────────────────────────────────────
try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except AttributeError:
    pass

current_time = datetime.now().strftime("%Y%m%d_%H%M")
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(levelname)-8s  %(message)s",
    handlers=[
        logging.FileHandler(f"facebook_scraper_log_{current_time}.txt", encoding="utf-8"),
        logging.StreamHandler(),
    ],
)
log = logging.getLogger(__name__)

# ─────────────────────────────────────────────────────────────────
# 1. CONFIG & API KEYS
# ─────────────────────────────────────────────────────────────────
# Mặc định cào từ trang Fanpage chính thức của Crossfire Legends
FB_PAGE_ID = "CrossfireLegendsVN"  # Hoặc ID dạng số của Page
OUTPUT_FILE = f"Crossfire_Legends_Facebook_Report_{current_time}.xlsx"

# ─────────────────────────────────────────────────────────────────
# 2. SENTIMENT & CATEGORIZATION (Đồng bộ với Store Scraper)
# ─────────────────────────────────────────────────────────────────
KEYWORD_DICT = {
    "vi": {
        "negative": [
            r"\blag\b", r"\bgiật\b", r"\bvăng\b", r"\blỗi\b", r"\btreo\b",
            r"\bhút máu\b", r"\brác\b", r"\bdis\b(?!cover|play|tance|count)",
            r"\btệ\b", r"\bchán\b", r"\bxấu\b", r"\bnát\b", r"\bxả rác\b",
            r"\bhack\b", r"\bcheat\b", r"\bhách\b", r"\bcit\b", r"\bhút tiền\b",
        ],
        "positive": [
            r"\bmượt\b", r"\bđỉnh\b", r"\bhay\b", r"\bcuốn\b", r"\bvui\b",
            r"\bthích\b", r"\btuyệt\b", r"\bổn\b", r"\bngon\b", r"\byêu\b",
        ],
    }
}

NEG_PATTERN = re.compile("|".join(KEYWORD_DICT["vi"]["negative"]), re.IGNORECASE)
POS_PATTERN = re.compile("|".join(KEYWORD_DICT["vi"]["positive"]), re.IGNORECASE)

def analyze_sentiment(text: str, rating: int = None) -> str:
    """
    Phân tích sắc thái của bình luận dựa trên từ khóa.
    Nếu có thông tin đánh giá (ví dụ bài review page), kết hợp thêm rating.
    """
    text_lower = str(text).lower()
    
    if rating is not None:
        if rating <= 2:
            return "Tiêu cực"
        if rating >= 4:
            if NEG_PATTERN.search(text_lower):
                return "Hỗn hợp"
            return "Tích cực"

    has_neg = bool(NEG_PATTERN.search(text_lower))
    has_pos = bool(POS_PATTERN.search(text_lower))
    
    if has_neg and has_pos:
        return "Hỗn hợp"
    if has_neg:
        return "Tiêu cực"
    if has_pos:
        return "Tích cực"
    return "Trung lập"

def classify_comment_multi(text: str, sentiment: str) -> list[str]:
    """
    Phân loại phản hồi của người chơi trên Facebook thành các danh mục lỗi/chủ đề cụ thể.
    """
    if sentiment == "Tích cực":
        return ["Khen"]
        
    text_lower = str(text).lower()
    matched = []
    
    # 1. Hack / Cheat
    if any(k in text_lower for k in ["hack", "cheat", "hách", "cit", "citer", "wallhack", "aimbot", "mod"]):
        matched.append("Vấn Nạn Hack / Cheat")
        
    # 2. Ghép Trận
    if any(k in text_lower for k in ["ghép trận", "ghep tran", "tìm trận", "tim tran", "ghép ranh", "ghep ranh", "ghép đội", "ghep doi", "tìm trận lâu"]):
        matched.append("Lỗi Ghép Trận")
        
    # 3. Kết Nối / Đăng Nhập
    if any(k in text_lower for k in ["kết nối", "ket noi", "đường truyền", "duong truyen", "đăng nhập", "dang nhap", "login", "server", "mạng yếu", "mất mạng", "wifi", "lỗi mạng", "discon"]):
        matched.append("Lỗi Kết Nối / Đăng Nhập")
        
    # 4. Văng Game / Crash
    if any(k in text_lower for k in ["văng", "vang", "crash", "out game", "out ranh", "bay ra", "bị off", "bị thoát", "thoát ra", "sập game", "đơ máy"]):
        matched.append("Lỗi Văng Game / Crash")
 
    # 5. Âm Thanh / Hình Ảnh
    if any(k in text_lower for k in ["mất tiếng", "mat tieng", "âm thanh", "loa", "mic", "hình ảnh", "hinh anh", "đồ họa"]):
        matched.append("Lỗi Âm Thanh / Hình Ảnh")
 
    # 6. Cập Nhật / Đứng Tải
    if any(k in text_lower for k in ["cập nhật", "cap nhat", "cập nhập", "đứng 100%", "đứng tải", "tải game", "không tải được", "load chậm", "ko vào được", "không vào được"]):
        matched.append("Lỗi Cập Nhật / Đứng Tải")
 
    # 7. Giật Lag / Drop FPS
    if any(k in text_lower for k in ["lag", "lác", "lắc", "giật", "giat", "ping", "fps", "stutter", "delay", "đơ"]):
        matched.append("Lỗi Giật Lag / Drop FPS")
        
    # 8. Nạp Tiền / Giao Dịch
    if any(k in text_lower for k in ["nạp", "nap", "topup", "hút máu", "hut mau", "kim cương", "nạp xu", "mất xu", "sự kiện nạp", "giá", "mua pass", "pay to win", "p2w"]):
        matched.append("Lỗi Nạp Tiền / Giao Dịch")
        
    # 9. Cân Bằng Gameplay
    if any(k in text_lower for k in ["zombie", "cân bằng", "can bang", "gameplay", "nerf", "buff", "súng", "sung", "dame", "lỗi map"]):
        matched.append("Cân Bằng Gameplay")
        
    if not matched:
        return ["Ý Kiến Khác / Chê Chung"]
        
    return matched

# ─────────────────────────────────────────────────────────────────
# 3. FACEBOOK SCRAPER CLIENT (SKELETON)
# ─────────────────────────────────────────────────────────────────
class FacebookScraper:
    def __init__(self, page_id: str, access_token: str = None):
        self.page_id = page_id
        self.access_token = access_token
        
    def fetch_posts(self, limit: int = 20) -> list[dict]:
        """
        Lấy danh sách các bài viết (posts) gần nhất từ Fanpage.
        Đây là hàm skeleton. Agent có thể chọn các giải pháp sau để implement:
        - Cách 1: Sử dụng Facebook Graph API (Yêu cầu Page Access Token)
        - Cách 2: Sử dụng các thư viện cào hoặc giả lập trình duyệt (Selenium / Playwright / HTTP Requests)
        """
        log.info(f"Đang lấy {limit} bài viết gần nhất từ Fanpage '{self.page_id}'...")
        
        # Mẫu dữ liệu giả định trả về từ Fanpage
        mock_posts = [
            {
                "post_id": "1000001",
                "message": "[CẬP NHẬT PHIÊN BẢN MỚI] Chiến binh ơi! Phiên bản mới đã chính thức cập bến. Hãy tải ngay để trải nghiệm chế độ Zombie v4 cực hot và loạt súng mới đỉnh cao!",
                "created_time": "2026-07-06T10:00:00+0700",
                "like_count": 150,
                "comment_count": 85
            },
            {
                "post_id": "1000002",
                "message": "Thông báo bảo trì định kỳ hệ thống máy chủ để nâng cấp đường truyền và tối ưu hóa tính năng ghép trận ranh rực lửa.",
                "created_time": "2026-07-05T23:30:00+0700",
                "like_count": 92,
                "comment_count": 41
            }
        ]
        
        # TODO: Cài đặt logic cào thực tế ở đây
        # Ví dụ nếu dùng Graph API:
        # url = f"https://graph.facebook.com/v19.0/{self.page_id}/posts"
        # params = {"access_token": self.access_token, "limit": limit, "fields": "id,message,created_time,shares,comments.summary(true)"}
        # res = requests.get(url, params=params) ...
        
        return mock_posts

    def fetch_comments(self, post_id: str, limit: int = 50) -> list[dict]:
        """
        Lấy các bình luận (comments) của một bài viết cụ thể.
        """
        log.info(f"Đang lấy bình luận của bài viết ID {post_id}...")
        
        # Mẫu dữ liệu bình luận giả định chứa phản hồi của người chơi
        mock_comments = {
            "1000001": [
                {"comment_id": "c1", "author": "Nguyễn Văn A", "message": "Game cập nhật xong lag quá admin ơi, bắn cứ giật giật bực cả mình!", "created_time": "2026-07-06T10:05:00+0700", "like_count": 12},
                {"comment_id": "c2", "author": "Trần Thị B", "message": "Bản mới mượt đỉnh chóp, súng mới bắn rất đã tay, chúc game ngày càng phát triển.", "created_time": "2026-07-06T10:08:00+0700", "like_count": 3},
                {"comment_id": "c3", "author": "Lê Văn C", "message": "Có ai bị lỗi không vào được game giống tôi không? Cứ đứng tải ở màn hình 100% mãi.", "created_time": "2026-07-06T10:15:00+0700", "like_count": 5},
                {"comment_id": "c4", "author": "Hoàng D", "message": "Súng zombie v4 dame lỗi hay sao ấy, bắn zombie mãi không chết, cần nerf lại bớt đi.", "created_time": "2026-07-06T10:20:00+0700", "like_count": 1},
                {"comment_id": "c5", "author": "Phạm E", "message": "Game toàn hack cheat đi ranh gặp suốt mà ko thấy admin khóa acc.", "created_time": "2026-07-06T10:22:00+0700", "like_count": 18}
            ],
            "1000002": [
                {"comment_id": "c6", "author": "Vũ F", "message": "Nâng cấp xong có sửa lỗi đăng nhập bằng Facebook không ad? Mấy hôm nay toàn báo lỗi kết nối mạng.", "created_time": "2026-07-05T23:35:00+0700", "like_count": 8},
                {"comment_id": "c7", "author": "Đặng G", "message": "Ghép trận ranh lâu kinh khủng, tìm cả buổi tối ko được trận nào.", "created_time": "2026-07-05T23:45:00+0700", "like_count": 10}
            ]
        }
        
        # TODO: Cài đặt logic cào thực tế ở đây
        
        return mock_comments.get(post_id, [])

# ─────────────────────────────────────────────────────────────────
# 4. EXCEL REPORT WRITER
# ─────────────────────────────────────────────────────────────────
def write_excel(df_comments: pd.DataFrame, df_posts: pd.DataFrame, filename: str):
    """
    Xuất dữ liệu đã xử lý ra file Excel định dạng đẹp mắt.
    """
    log.info(f"Đang ghi dữ liệu vào file Excel: {filename}...")
    
    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        df_comments.to_excel(writer, sheet_name="FB_Comments_Raw", index=False)
        df_posts.to_excel(writer, sheet_name="FB_Posts_Summary", index=False)
        
        # Styling sheet chính
        workbook = writer.book
        sheet = workbook["FB_Comments_Raw"]
        
        # Thiết lập header format
        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        
        for col_num in range(1, len(df_comments.columns) + 1):
            cell = sheet.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            
        # Tự động căn chỉnh độ rộng cột
        for col in sheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 50)
            
    log.info("✓ Xuất file Excel hoàn thành thành công.")

# ─────────────────────────────────────────────────────────────────
# 5. MAIN PIPELINE
# ─────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    log.info("=" * 60)
    log.info("  CFL Facebook Scraper & Analyzer STARTED")
    log.info("=" * 60)
    
    # 1. Khởi tạo client cào
    scraper = FacebookScraper(page_id=FB_PAGE_ID)
    
    # 2. Cào bài viết
    posts = scraper.fetch_posts(limit=5)
    df_posts = pd.DataFrame(posts)
    
    # 3. Duyệt qua các bài viết để cào bình luận
    all_comments = []
    for post in posts:
        comments = scraper.fetch_comments(post["post_id"], limit=50)
        for c in comments:
            # Liên kết comment với bài viết nguồn
            c["source_post_id"] = post["post_id"]
            all_comments.append(c)
            
    if not all_comments:
        log.warning("❌ Không tìm thấy bình luận nào để xử lý.")
        sys.exit(0)
        
    df_comments = pd.DataFrame(all_comments)
    
    # 4. Phân tích sắc thái (Sentiment) và phân loại danh mục lỗi
    log.info("Đang thực hiện phân tích Sentiment và phân loại danh mục lỗi...")
    df_comments["sentiment"] = df_comments["message"].apply(analyze_sentiment)
    
    categories_list = []
    for idx, row in df_comments.iterrows():
        cats = classify_comment_multi(row["message"], row["sentiment"])
        categories_list.append(", ".join(cats))
        
    df_comments["categories"] = categories_list
    
    # 5. Xuất kết quả
    write_excel(df_comments, df_posts, OUTPUT_FILE)
    
    log.info("=" * 60)
    log.info(f"  TỔNG SỐ BÀI VIẾT QUÉT   : {len(df_posts)}")
    log.info(f"  TỔNG SỐ BÌNH LUẬN QUÉT  : {len(df_comments)}")
    log.info(f"  FILE KẾT QUẢ            : {OUTPUT_FILE}")
    log.info("=" * 60)
