"""LLM-synthesized insight summary + Excel export."""
from __future__ import annotations

import io
from datetime import datetime

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from ..database import get_db
from ..models import Analysis, Comment
from ..routers.stats import overview as _overview
from ..services.insights import generate_summary
from ..taxonomy import SENTIMENT_LABELS_VI, TOPIC_LABELS_VI

router = APIRouter(prefix="/api", tags=["insights"])


@router.get("/insights/summary")
def insights_summary(
    db: Session = Depends(get_db),
    source: str | None = None,
    date_from: datetime | None = Query(None, alias="from"),
    date_to: datetime | None = Query(None, alias="to"),
):
    ov = _overview(db=db, source=source, date_from=date_from, date_to=date_to)

    stmt = (
        select(Comment.message)
        .join(Analysis, Analysis.comment_id == Comment.id)
        .where(Analysis.sentiment == "negative")
        .order_by(Comment.id.desc())
        .limit(8)
    )
    if source:
        stmt = stmt.where(Comment.source_type == source)
    if date_from:
        stmt = stmt.where(Comment.created_at >= date_from)
    if date_to:
        stmt = stmt.where(Comment.created_at <= date_to)
    samples = [m for (m,) in db.execute(stmt)]

    cache_key = (source, str(date_from), str(date_to), ov["total_comments"], ov["analyzed"])
    text = generate_summary(ov, samples, cache_key)
    return {"summary": text, "based_on": ov}


@router.get("/export")
def export_excel(
    db: Session = Depends(get_db),
    source: str | None = None,
    topic: str | None = None,
    sentiment: str | None = None,
    urgency: str | None = None,
    date_from: datetime | None = Query(None, alias="from"),
    date_to: datetime | None = Query(None, alias="to"),
):
    stmt = (
        select(Comment)
        .options(selectinload(Comment.analysis))
        .join(Analysis, Analysis.comment_id == Comment.id, isouter=True)
    )
    if source:
        stmt = stmt.where(Comment.source_type == source)
    if topic:
        stmt = stmt.where(Analysis.topic_main == topic)
    if sentiment:
        stmt = stmt.where(Analysis.sentiment == sentiment)
    if urgency:
        stmt = stmt.where(Analysis.urgency == urgency)
    if date_from:
        stmt = stmt.where(Comment.created_at >= date_from)
    if date_to:
        stmt = stmt.where(Comment.created_at <= date_to)
    stmt = stmt.order_by(Comment.created_at.desc().nullslast())

    rows = list(db.scalars(stmt).unique())

    wb = Workbook()
    ws = wb.active
    ws.title = "Feedback"
    headers = ["ID", "Nguồn", "Thời gian", "Nội dung", "Rating", "Chủ đề chính",
              "Chủ đề phụ", "Sentiment", "Mức độ khẩn cấp", "Tóm tắt", "Độ tin cậy"]
    ws.append(headers)
    header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill

    for c in rows:
        a = c.analysis
        ws.append([
            c.id,
            c.source_type,
            c.created_at.strftime("%Y-%m-%d %H:%M") if c.created_at else "",
            c.message,
            c.rating or "",
            TOPIC_LABELS_VI.get(a.topic_main, a.topic_main) if a else "",
            ", ".join(TOPIC_LABELS_VI.get(t, t) for t in (a.topics_sub or [])) if a else "",
            SENTIMENT_LABELS_VI.get(a.sentiment, a.sentiment) if a else "",
            a.urgency if a else "",
            a.summary if a else "",
            round(a.confidence, 2) if a else "",
        ])

    for col_idx, header in enumerate(headers, start=1):
        width = 60 if header in ("Nội dung", "Tóm tắt") else 16
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"CFL_Feedback_Export_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
