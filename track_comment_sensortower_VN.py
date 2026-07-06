"""
  Crossfire Legends - Review Scraper via Sensor Tower API (Vietnam Only)
  Platforms : Google Play + App Store (VN region)
  Date Range: 2026-03-03 to Present (Version 2 & Version 3 segments)
"""

import os
import re
import sys
import time
import logging
import requests
import pandas as pd
from datetime import datetime, timedelta
from pathlib import Path
from dotenv import load_dotenv
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ─────────────────────────────────────────────────────────────────
# 0. LOGGING & STREAM ENCODING
# ─────────────────────────────────────────────────────────────────
try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except AttributeError:
    pass

current_time = datetime.now().strftime("%Y%m%d_%H%M")
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(levelname)-8s  %(message)s",
    handlers=[
        logging.FileHandler(f"sensortower_vn_scraper_log_{current_time}.txt", encoding="utf-8"),
        logging.StreamHandler(),
    ],
)
log = logging.getLogger(__name__)

# ─────────────────────────────────────────────────────────────────
# 1. CONFIG & API KEYS
# ─────────────────────────────────────────────────────────────────
# Load env files
load_dotenv("key_api.env")
SENSORTOWER_API_KEY = os.getenv("SENSORTOWER_API_KEY")

OUTPUT_FILE = f"Crossfire_Legends_SensorTower_VN_Report_{current_time}.xlsx"

# App configurations (App IDs/Package Names in Sensor Tower) - VIETNAM ONLY
GOOGLE_APPS = {
    "VN": "com.vnggames.cfl.crossfirelegends",
}

APPLE_APPS = {
    "VN": "6748588650",
}

# Start date for reviews
START_DATE = datetime(2026, 3, 3)
END_DATE = datetime.now()

# ─────────────────────────────────────────────────────────────────
# 2. SENTIMENT & CATEGORIZATION
# ─────────────────────────────────────────────────────────────────
KEYWORD_DICT = {
    "en": {
        "negative": [
            r"\blag\b", r"\blagging\b", r"\bcrash(es|ed)?\b", r"\bbug(s|gy)?\b",
            r"\bdisconnect(ed|s)?\b", r"\bcheat(er|s)?\b", r"\bhack(er|s)?\b",
            r"\bpay[\s\-]?to[\s\-]?win\b", r"\bp2w\b", r"\bworst\b", r"\bterrible\b",
            r"\bawful\b", r"\bwaste\b", r"\bscam\b", r"\bunplayable\b",
            r"\bfreeze(s|d)?\b", r"\bglitch(es|y)?\b", r"\bstutter(s|ing)?\b",
            r"\bhigh ping\b", r"\bserver (down|issue|error)\b",
        ],
        "positive": [
            r"\bsmooth\b", r"\bgood\b", r"\bgreat\b", r"\bawesome\b",
            r"\blove\b", r"\bnice\b", r"\bbest\b", r"\bfun\b", r"\benjoy\b",
            r"\bperfect\b", r"\bexcellent\b", r"\bamazing\b", r"\bfantastic\b",
            r"\bstable\b", r"\bno lag\b", r"\bno bug\b",
        ],
    },
    "vi": {
        "negative": [
            r"\blag\b", r"\bgiật\b", r"\bvăng\b", r"\blỗi\b", r"\btreo\b",
            r"\bhút máu\b", r"\brác\b", r"\bdis\b(?!cover|play|tance|count)",
            r"\btệ\b", r"\bchán\b", r"\bxấu\b", r"\bnát\b", r"\bxả rác\b",
        ],
        "positive": [
            r"\bmượt\b", r"\bđỉnh\b", r"\bhay\b", r"\bcuốn\b", r"\bvui\b",
            r"\bthích\b", r"\btuyệt\b", r"\bổn\b", r"\bngon\b",
        ],
    },
}

NEG_PATTERN = re.compile("|".join(
    token for lang_dict in KEYWORD_DICT.values() for token in lang_dict["negative"]
), re.IGNORECASE)

POS_PATTERN = re.compile("|".join(
    token for lang_dict in KEYWORD_DICT.values() for token in lang_dict["positive"]
), re.IGNORECASE)

def analyze_sentiment(row: pd.Series) -> str:
    rating = int(row.get("rating", 3))
    text   = str(row.get("content", ""))

    if rating <= 2:
        return "Tiêu cực"
    if rating >= 4:
        if NEG_PATTERN.search(text):
            return "Hỗn hợp"
        return "Tích cực"

    has_neg = bool(NEG_PATTERN.search(text))
    has_pos = bool(POS_PATTERN.search(text))
    if has_neg and has_pos:
        return "Hỗn hợp"
    if has_neg:
        return "Tiêu cực"
    if has_pos:
        return "Tích cực"
    return "Trung lập"

CATEGORY_SEVERITY_ORDER = [
    "Vấn Nạn Hack / Cheat",
    "Lỗi Văng Game / Crash",
    "Lỗi Giật Lag / Drop FPS",
    "Lỗi Kết Nối / Đăng Nhập",
    "Lỗi Ghép Trận",
    "Lỗi Cập Nhật / Đứng Tải",
    "Lỗi Âm Thanh / Hình Ảnh",
    "Lỗi Nạp Tiền / Giao Dịch",
    "Cân Bằng Gameplay"
]

def classify_comment_multi(text, sentiment):
    if sentiment == "Tích cực":
        return ["Khen"]
        
    text_lower = str(text).lower()
    matched = []
    
    # 1. Cheating / Hacking (Vấn Nạn Hack / Cheat)
    if any(k in text_lower for k in ["hack", "cheat", "hách", "cit", "citer", "wallhack", "aimbot", "mod"]):
        matched.append("Vấn Nạn Hack / Cheat")
        
    # 2. Matchmaking Issue (Lỗi Ghép Trận)
    if any(k in text_lower for k in ["ghép trận", "ghep tran", "tìm trận", "tim tran", "ghép ranh", "ghep ranh", "ghéo trận", "ghép đội", "ghep doi", "chọn map", "chọn màn", "chọn bản đồ", "kẹt ở map", "tìm cả ngày", "ghép lâu", "không tìm thấy người chơi", "không tìm thấy người", "tìm trận lâu"]):
        matched.append("Lỗi Ghép Trận")
        
    # 3. Connection / Login / Server Issue (Lỗi Kết Nối / Đăng Nhập)
    if any(k in text_lower for k in ["kết nối", "ket noi", "đường truyền", "duong truyen", "đăng nhập", "dang nhap", "login", "server", "mạng yếu", "mất mạng", "wifi", "lỗi mạng", "văng phòng", "discon", "đăng nhập game", "mất kết nối", "lỗi đường truyền", "loi duong truyen"]):
        matched.append("Lỗi Kết Nối / Đăng Nhập")
        
    # 4. Crash / Force Close (Lỗi Văng Game / Crash)
    if any(k in text_lower for k in ["văng", "vang", "crash", "out game", "out ranh", "bay ra", "bị off", "bi off", "bị thoát", "thoát ra", "sập game", "sap game", "tự đóng", "tu dong", "đơ máy", "do may"]):
        matched.append("Lỗi Văng Game / Crash")

    # 5. Audio / Visual Issue (Lỗi Âm Thanh / Hình Ảnh)
    if any(k in text_lower for k in ["mất tiếng", "mat tieng", "âm thanh", "am thanh", "tiếng game", "loa", "mic", "hình ảnh", "hinh anh", "chế độ zoom", "lỗi zoom", "lỗi hình", "đồ họa", "giọng"]):
        matched.append("Lỗi Âm Thanh / Hình Ảnh")

    # 6. Update / Loading Issue (Lỗi Cập Nhật / Đứng Tải)
    if any(k in text_lower for k in ["cập nhật", "cap nhat", "cập nhập", "cap nhap", "đứng 100%", "đứng tải", "tải tài nguyên", "tải game", "cài game", "không tải được", "khong tai duoc", "load chậm", "kẹt màn hình", "treo màn hình", "vào trận ko", "vô trận ko", "chơi mà không vô được", "vào ko vào được", "ko vào được", "vào không được", "vào ko được", "ko choi dc", "vào ko choi dc", "không vào được", "cập nhận"]):
        matched.append("Lỗi Cập Nhật / Đứng Tải")

    # 7. Lag / Ping / FPS Issue (Lỗi Giật Lag / Drop FPS)
    if any(k in text_lower for k in ["lag", "lác", "lắc", "giật", "giat", "ping", "fps", "stutter", "delay", "đơ", "đứng hình", "đứng im", "giật lắt"]):
        matched.append("Lỗi Giật Lag / Drop FPS")
        
    # 8. Top-up / Monetization (Lỗi Nạp Tiền / Giao Dịch)
    if any(k in text_lower for k in ["nạp", "nap", "topup", "hút máu", "hut mau", "mảnh đổi", "manh doi", "kim cương", "kim cuong", "nạp xu", "mất xu", "đầu tư", "tiền", "tien", "sự kiện nạp", "giá", "mở rương", "quay rương", "nạp pass", "mua pass", "vé pass", "mảnh skin", "vip bạc", "free", "phải nạp", "bắt nạp", "pay to win", "p2w"]):
        matched.append("Lỗi Nạp Tiền / Giao Dịch")
        
    # 9. Gameplay Balance (Cân Bằng Gameplay)
    if any(k in text_lower for k in ["zombie", "cân bằng", "can bang", "gameplay", "nerf", "buff", "súng", "sung", "bắn trúng", "dame", "nhảy cao", "lỗi map", "zombi", "trừ điểm", "cộng điểm", "đội thua", "mạnh quá"]):
        matched.append("Cân Bằng Gameplay")
        
    if not matched:
        return ["Ý Kiến Khác / Chê Chung"]
        
    return matched

def categorize_feedback(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    if df.empty:
        df["category"] = []
        df["all_categories"] = []
        return df
        
    categories = []
    all_categories = []
    
    for idx, row in df.iterrows():
        matched = classify_comment_multi(row["content"], row["sentiment"])
        categories.append(", ".join(matched))
        all_categories.append(", ".join(matched))
        
    df["category"] = categories
    df["all_categories"] = all_categories
    return df

# ─────────────────────────────────────────────────────────────────
# 3. SENSOR TOWER API CLIENT
# ─────────────────────────────────────────────────────────────────
def fetch_sensortower_reviews(os_platform: str, app_id: str, region_code: str) -> pd.DataFrame:
    """Fetch reviews from Sensor Tower endpoint"""
    if not SENSORTOWER_API_KEY or SENSORTOWER_API_KEY == "YOUR_SENSORTOWER_API_KEY_HERE":
        log.error("API Key for Sensor Tower is missing! Please set SENSORTOWER_API_KEY in key_api.env.")
        return pd.DataFrame()
        
    all_reviews = []
    store_name = "Google Play" if os_platform == "android" else "App Store"
    log.info(f"[{region_code}] Fetching reviews for {store_name} ({app_id}) from Sensor Tower...")
    
    page = 1
    limit = 200
    start_date_str = START_DATE.strftime("%Y-%m-%d")
    end_date_str = END_DATE.strftime("%Y-%m-%d")
    
    url = f"https://api.sensortower.com/v1/{os_platform}/review/get_reviews"
    
    while True:
        params = {
            "auth_token": SENSORTOWER_API_KEY,
            "app_id": app_id,
            "country": region_code,
            "start_date": start_date_str,
            "end_date": end_date_str,
            "limit": limit,
            "page": page
        }
        
        try:
            res = requests.get(url, params=params, timeout=20)
            if res.status_code == 401:
                log.error("401 Unauthorized: Invalid Sensor Tower API key.")
                return pd.DataFrame()
            elif res.status_code == 403:
                log.error("403 Forbidden: No permissions to reviews API for this account package.")
                return pd.DataFrame()
            elif res.status_code != 200:
                log.warning(f"HTTP Error {res.status_code} for page {page}.")
                break
                
            data = res.json()
            
            # Handle list or dict containing a list
            if isinstance(data, list):
                reviews_list = data
            elif isinstance(data, dict):
                reviews_list = data.get("feedback", data.get("reviews", data.get("data", [])))
                # Some API designs return reviews directly under the top-level keys if it's formatted differently
                if not isinstance(reviews_list, list):
                    reviews_list = []
                    for v in data.values():
                        if isinstance(v, list):
                            reviews_list = v
                            break
            else:
                reviews_list = []
                
            if not reviews_list:
                log.info(f"  No more reviews returned at page {page}.")
                break
                
            page_fetched = 0
            for r in reviews_list:
                # Robust key extraction
                body = r.get("body", r.get("content", r.get("text", r.get("review", ""))))
                title = r.get("title", "")
                full_content = body
                if title and title != body:
                    full_content = f"{title} - {body}"
                    
                author = r.get("username", r.get("author", r.get("reviewer", r.get("nickname", "User"))))
                rating = r.get("rating", r.get("star_rating", r.get("score", 3)))
                try:
                    rating = int(rating)
                except:
                    rating = 3
                    
                date_val = r.get("date", r.get("updated_at", r.get("created_at", r.get("at", ""))))
                
                all_reviews.append({
                    "author": author,
                    "rating": rating,
                    "content": full_content,
                    "at": date_val,
                    "Region": region_code,
                    "Store": store_name,
                })
                page_fetched += 1
                
            log.info(f"  Page {page}: fetched {page_fetched} reviews from Sensor Tower.")
            
            # If the response list is smaller than the limit, it means we reached the last page
            if page_fetched < limit:
                break
                
            page += 1
            time.sleep(0.5)
        except Exception as e:
            log.error(f"Error fetching page {page} from Sensor Tower: {e}")
            break
            
    df = pd.DataFrame(all_reviews)
    log.info(f"[{region_code}] Completed {store_name}. Total reviews fetched: {len(df)}")
    return df

# ─────────────────────────────────────────────────────────────────
# 4. REPORT STYLING & GENERATION
# ─────────────────────────────────────────────────────────────────
def assign_version_segment(dt):
    if pd.isnull(dt):
        return "Unknown"
    if dt.tzinfo is not None:
        dt = dt.tz_localize(None)
    
    v2_start = datetime(2026, 3, 3)
    v2_end = datetime(2026, 4, 27, 23, 59, 59)
    
    if dt < v2_start:
        return "Trước V2"
    elif dt <= v2_end:
        return "Version 2 (03/03 - 27/04)"
    else:
        return "Version 3 (28/04 - Present)"

def get_store_version_stats(df, store_name, segment_name):
    sub_df = df[(df["Store"] == store_name) & (df["Version Segment"] == segment_name)]
    total = len(sub_df)
    avg_rating = sub_df["rating"].mean() if total > 0 else 0.0
    pos_count = len(sub_df[sub_df["sentiment"] == "Tích cực"])
    neg_count = len(sub_df[sub_df["sentiment"] == "Tiêu cực"])
    
    pos_pct = (pos_count / total * 100) if total > 0 else 0.0
    neg_pct = (neg_count / total * 100) if total > 0 else 0.0
    
    neg_df = sub_df[sub_df["sentiment"].isin(["Tiêu cực", "Hỗn hợp"])]
    if not neg_df.empty:
        temp_neg = neg_df.copy()
        temp_neg["category"] = temp_neg["category"].str.split(", ")
        exploded_neg = temp_neg.explode("category")
        top_cats = exploded_neg["category"].value_counts()
        top_cats_clean = top_cats.drop("Ý Kiến Khác / Chê Chung", errors="ignore")
        top_issue = top_cats_clean.index[0] if not top_cats_clean.empty else (top_cats.index[0] if not top_cats.empty else "Không có")
    else:
        top_issue = "Không có"
        
    return {
        "total": total,
        "avg_rating": avg_rating,
        "pos_pct": pos_pct,
        "neg_pct": neg_pct,
        "pos_count": pos_count,
        "neg_count": neg_count,
        "top_issue": top_issue,
        "neg_df": neg_df
    }

def generate_platform_text_analysis(v2_stats, v3_stats, store_name):
    total_v2 = v2_stats["total"]
    total_v3 = v3_stats["total"]
    avg_v2 = v2_stats["avg_rating"]
    avg_v3 = v3_stats["avg_rating"]
    neg_pct_v2 = v2_stats["neg_pct"]
    neg_pct_v3 = v3_stats["neg_pct"]
    top_issue_v2 = v2_stats["top_issue"]
    top_issue_v3 = v3_stats["top_issue"]
    
    if total_v2 == 0 and total_v3 == 0:
        return f"Không có dữ liệu đánh giá trên {store_name} trong cả hai phiên bản."
        
    analysis = f"Phân tích đánh giá trên {store_name}:\\n"
    if total_v2 > 0 and total_v3 > 0:
        analysis += f"• Quy mô dữ liệu: Version 2 ghi nhận {total_v2} đánh giá, trong khi Version 3 ghi nhận {total_v3} đánh giá.\\n"
        
        # Rating trend
        rating_diff = avg_v3 - avg_v2
        if rating_diff > 0.05:
            analysis += f"• Xu hướng điểm số: Điểm đánh giá trung bình tăng từ {avg_v2:.2f} (V2) lên {avg_v3:.2f} (V3), cho thấy tín hiệu cải thiện tích cực về mức độ hài lòng chung của người chơi.\\n"
        elif rating_diff < -0.05:
            analysis += f"• Xu hướng điểm số: Điểm đánh giá trung bình bị giảm sút đáng kể từ {avg_v2:.2f} (V2) xuống còn {avg_v3:.2f} (V3). Điều này cảnh báo phiên bản mới đang phát sinh thêm lỗi gây khó chịu cho người dùng.\\n"
        else:
            analysis += f"• Xu hướng điểm số: Điểm đánh giá trung bình giữ mức ổn định từ {avg_v2:.2f} (V2) sang {avg_v3:.2f} (V3) (chênh lệch {rating_diff:.2f}).\\n"
            
        # Sentiment trend
        neg_diff = neg_pct_v3 - neg_pct_v2
        if neg_diff > 3.0:
            analysis += f"• Tỷ lệ tiêu cực: Tỷ lệ phản hồi tiêu cực tăng thêm {neg_diff:.1f}% ở Version 3 (từ {neg_pct_v2:.1f}% lên {neg_pct_v3:.1f}%), chứng tỏ phiên bản mới cần được tối ưu hóa kỹ lưỡng hơn.\\n"
        elif neg_diff < -3.0:
            analysis += f"• Tỷ lệ tiêu cực: Tỷ lệ phản hồi tiêu cực giảm rõ rệt {abs(neg_diff):.1f}% ở Version 3 (từ {neg_pct_v2:.1f}% xuống {neg_pct_v3:.1f}%), cho thấy những nỗ lực cập nhật/vá lỗi đã phát huy tác dụng.\\n"
        else:
            analysis += f"• Tỷ lệ tiêu cực: Tỷ lệ tiêu cực biến động nhẹ (V2: {neg_pct_v2:.1f}% so với V3: {neg_pct_v3:.1f}%).\\n"
            
        # Issue category comparison
        if top_issue_v2 == top_issue_v3:
            if top_issue_v3 != "Không có" and top_issue_v3 != "Khen":
                analysis += f"• Vấn đề cốt lõi: Nhóm '{top_issue_v3}' tiếp tục là vấn đề nghiêm trọng nhất ở cả 2 phiên bản. Đây là điểm nghẽn trải nghiệm chưa được giải quyết triệt để.\\n"
            else:
                analysis += f"• Vấn đề cốt lõi: Không ghi nhận vấn đề kỹ thuật đặc thù nào quá nổi trội.\\n"
        else:
            analysis += f"• Sự thay đổi lỗi trọng tâm: Ở Version 2, vấn đề nổi cộm nhất là '{top_issue_v2}', nhưng sang Version 3 đã chuyển dịch sang '{top_issue_v3}'. Vận hành cần tập trung xử lý ngay vấn đề mới phát sinh này.\\n"
    elif total_v3 > 0:
        analysis += f"• Chỉ ghi nhận dữ liệu cho Version 3 với {total_v3} đánh giá, điểm TB là {avg_v3:.2f}/5 sao. Tỷ lệ tiêu cực chiếm {neg_pct_v3:.1f}%. Vấn đề nổi bật nhất là '{top_issue_v3}'.\\n"
    else:
        analysis += f"• Chỉ ghi nhận dữ liệu cho Version 2 với {total_v2} đánh giá, điểm TB là {avg_v2:.2f}/5 sao. Tỷ lệ tiêu cực chiếm {neg_pct_v2:.1f}%. Vấn đề nổi bật nhất là '{top_issue_v2}'.\\n"
        
    return analysis

def generate_recommendations(gp_v3_stats, as_v3_stats):
    recommendations = ["Dựa trên phản hồi thực tế của người dùng đối với phiên bản mới nhất (Version 3) từ ngày 28/04:\\n"]
    
    gp_issue = gp_v3_stats["top_issue"]
    as_issue = as_v3_stats["top_issue"]
    
    # Count top issues in V3 across both platforms
    all_v3_neg_cats = []
    if not gp_v3_stats["neg_df"].empty:
        all_v3_neg_cats.append(gp_v3_stats["neg_df"])
    if not as_v3_stats["neg_df"].empty:
        all_v3_neg_cats.append(as_v3_stats["neg_df"])
        
    if all_v3_neg_cats:
        v3_neg_df = pd.concat(all_v3_neg_cats, ignore_index=True).copy()
        v3_neg_df["category"] = v3_neg_df["category"].str.split(", ")
        exploded_neg = v3_neg_df.explode("category")
        v3_cats = exploded_neg["category"].value_counts().drop("Ý Kiến Khác / Chê Chung", errors="ignore")
    else:
        v3_cats = pd.Series()
        
    count = 1
    if not v3_cats.empty:
        top_v3_issue = v3_cats.index[0]
        recommendations.append(f"{count}. Tập trung khắc phục triệt để lỗi thuộc nhóm '{top_v3_issue}' trên cả 2 nền tảng, đây là nguyên nhân hàng đầu gây sụt giảm rating của Version 3.")
        count += 1
        
        # Specific recommendations depending on what issues appear in Version 3
        if "Vấn Nạn Hack / Cheat" in v3_cats.index:
            recommendations.append(f"{count}. [Bảo mật] Tăng cường hệ thống quét hack/cheat tự động cho Version 3, phản hồi nhanh các tố cáo gian lận của người dùng.")
            count += 1
        if "Lỗi Văng Game / Crash" in v3_cats.index:
            recommendations.append(f"{count}. [Tương thích] Kiểm tra log crash và tối ưu hóa tài nguyên cài đặt của Version 3 nhằm giảm thiểu văng game (crash) đột ngột.")
            count += 1
        if "Lỗi Giật Lag / Drop FPS" in v3_cats.index:
            recommendations.append(f"{count}. [Hiệu năng] Tối ưu hóa FPS và ping máy chủ cho Version 3, đặc biệt là vào các khung giờ cao điểm để giảm giật lag.")
            count += 1
        if "Lỗi Nạp Tiền / Giao Dịch" in v3_cats.index:
            recommendations.append(f"{count}. [Giao dịch] Rà soát lại luồng nạp thẻ và giao dịch của phiên bản mới để tránh chậm trễ chuyển vật phẩm/tiền ảo cho người chơi.")
            count += 1
        if "Lỗi Kết Nối / Đăng Nhập" in v3_cats.index:
            recommendations.append(f"{count}. [Máy chủ] Khắc phục tình trạng mất kết nối giữa trận và lỗi tải tài nguyên khi bắt đầu đăng nhập game trên Version 3.")
            count += 1
    else:
        recommendations.append(f"{count}. Tiếp tục duy trì chất lượng dịch vụ hiện tại, theo dõi sát các phản hồi của người dùng để phản ứng nhanh với bất kỳ vấn đề phát sinh nào.")
        
    return "\\n".join(recommendations)

def write_comparison_table(ws, start_row, title_prefix, v2_stats, v3_stats, fill_color="2C3E50"):
    c1_hdr = ws.cell(row=start_row, column=1, value=f"Chỉ số ({title_prefix})")
    c2_hdr = ws.cell(row=start_row, column=2, value="Version 2 (03/03 - 27/04)")
    c3_hdr = ws.cell(row=start_row, column=3, value="Version 3 (28/04 - Present)")
    
    font_hdr = Font(bold=True, color="FFFFFF")
    fill_hdr = PatternFill("solid", fgColor=fill_color)
    align_hdr = Alignment(horizontal="center", vertical="center")
    
    for cell in [c1_hdr, c2_hdr, c3_hdr]:
        cell.font = font_hdr
        cell.fill = fill_hdr
        cell.alignment = align_hdr
        
    metrics = [
        ("Tổng số đánh giá (Total Reviews)", v2_stats["total"], v3_stats["total"]),
        ("Điểm đánh giá TB (Avg Rating)", round(v2_stats["avg_rating"], 2), round(v3_stats["avg_rating"], 2)),
        ("Tỷ lệ Tích cực (Positive %)", f"{v2_stats['pos_pct']:.1f}% ({v2_stats['pos_count']} cmt)", f"{v3_stats['pos_pct']:.1f}% ({v3_stats['pos_count']} cmt)"),
        ("Tỷ lệ Tiêu cực (Negative %)", f"{v2_stats['neg_pct']:.1f}% ({v2_stats['neg_count']} cmt)", f"{v3_stats['neg_pct']:.1f}% ({v3_stats['neg_count']} cmt)"),
        ("Vấn đề chính (Top Issue)", v2_stats["top_issue"], v3_stats["top_issue"]),
    ]
    
    thin = Side(border_style="thin", color="D3D3D3")
    border = Border(top=thin, left=thin, right=thin, bottom=thin)
    
    for idx, (label, v2_val, v3_val) in enumerate(metrics, start_row + 1):
        c1 = ws.cell(row=idx, column=1, value=label)
        c2 = ws.cell(row=idx, column=2, value=v2_val)
        c3 = ws.cell(row=idx, column=3, value=v3_val)
        
        c1.alignment = Alignment(horizontal="left", vertical="center")
        c2.alignment = Alignment(horizontal="center", vertical="center")
        c3.alignment = Alignment(horizontal="center", vertical="center")
        
        for c in [c1, c2, c3]:
            c.border = border
            if idx % 2 == 1:
                c.fill = PatternFill("solid", fgColor="F2F4F4")

SENTIMENT_COLORS = {
    "Tiêu cực": "FFCCCC", "Tích cực": "CCFFCC",
    "Hỗn hợp": "FFF2CC", "Trung lập": "E0E0E0",
}

def _auto_column_width(ws):
    for col in ws.columns:
        max_len = max((len(str(cell.value)) for cell in col if cell.value), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 60)

def _style_header(ws, fill_hex="2C3E50"):
    fill = PatternFill("solid", fgColor=fill_hex)
    font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")

def style_raw_sheet(ws, sentiment_col_idx: int):
    _style_header(ws)
    for row in ws.iter_rows(min_row=2):
        sentiment_val = row[sentiment_col_idx - 1].value
        hex_color = SENTIMENT_COLORS.get(str(sentiment_val), "FFFFFF")
        fill = PatternFill("solid", fgColor=hex_color)
        for cell in row:
            cell.fill = fill

def write_excel(master_df: pd.DataFrame, bad_df: pd.DataFrame, top_keywords_df: pd.DataFrame, output_path: str):
    if master_df.empty:
        log.error("Master DataFrame is empty. No excel report generated.")
        return
        
    summary_pivot = pd.pivot_table(
        master_df, values="content", index="Version Segment", columns="Store",
        aggfunc="count", fill_value=0
    ).reset_index()
    store_cols = [c for c in summary_pivot.columns if c != "Version Segment"]
    summary_pivot["Total"] = summary_pivot[store_cols].sum(axis=1)

    sentiment_dist = (
        master_df.groupby(["Version Segment", "Store", "sentiment"]).size()
        .unstack(fill_value=0).reset_index()
    )

    rating_dist = (
        master_df.groupby(["Version Segment", "Store", "rating"])
        .size().reset_index(name="count")
    )

    # Explode categories for category_dist
    temp_cats_df = master_df[master_df["sentiment"].isin(["Tiêu cực", "Hỗn hợp"])].copy()
    temp_cats_df["category"] = temp_cats_df["category"].str.split(", ")
    exploded_df = temp_cats_df.explode("category")
    category_dist = (
        exploded_df.groupby(["Version Segment", "Store", "category"]).size()
        .reset_index(name="count")
        .sort_values(["Version Segment", "Store", "count"], ascending=[True, True, False])
    )

    gp_v2_stats = get_store_version_stats(master_df, "Google Play", "Version 2 (03/03 - 27/04)")
    gp_v3_stats = get_store_version_stats(master_df, "Google Play", "Version 3 (28/04 - Present)")
    as_v2_stats = get_store_version_stats(master_df, "App Store", "Version 2 (03/03 - 27/04)")
    as_v3_stats = get_store_version_stats(master_df, "App Store", "Version 3 (28/04 - Present)")

    gp_text = generate_platform_text_analysis(gp_v2_stats, gp_v3_stats, "Google Play")
    as_text = generate_platform_text_analysis(as_v2_stats, as_v3_stats, "App Store")
    rec_text = generate_recommendations(gp_v3_stats, as_v3_stats)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        wb = writer.book
        ws_ov = wb.create_sheet(title="Overview_&_Assessment")
        
        # Title Blocks
        ws_ov.merge_cells("A1:F1")
        ws_ov["A1"] = "BÁO CÁO TỔNG QUAN & ĐÁNH GIÁ TRẢI NGHIỆM NGƯỜI DÙNG VN (SENSOR TOWER)"
        ws_ov["A1"].font = Font(bold=True, size=14, color="1F4E79")
        ws_ov["A1"].alignment = Alignment(horizontal="center", vertical="center")
        
        ws_ov.merge_cells("A2:F2")
        ws_ov["A2"] = f"Dữ liệu từ {START_DATE.strftime('%d/%m/%Y')} đến {END_DATE.strftime('%d/%m/%Y')} | Phân khúc V2 vs V3"
        ws_ov["A2"].font = Font(italic=True, size=10, color="595959")
        ws_ov["A2"].alignment = Alignment(horizontal="center", vertical="center")
        
        # Google Play Section
        ws_ov["A4"] = "1. THỐNG KÊ & SO SÁNH PHIÊN BẢN CỬA HÀNG GOOGLE PLAY"
        ws_ov["A4"].font = Font(bold=True, size=11, color="1F4E79")
        
        write_comparison_table(ws_ov, 5, "Google Play", gp_v2_stats, gp_v3_stats, "1F4E79")
        
        ws_ov["A12"] = "PHÂN TÍCH & ĐÁNH GIÁ CHI TIẾT GOOGLE PLAY"
        ws_ov["A12"].font = Font(bold=True, size=11, color="2C3E50")
        
        ws_ov.merge_cells("A13:F15")
        ws_ov["A13"] = gp_text
        ws_ov["A13"].alignment = Alignment(wrap_text=True, vertical="top")
        ws_ov["A13"].font = Font(size=10)
        ws_ov["A13"].fill = PatternFill("solid", fgColor="EBF5FB")
        
        # App Store Section
        ws_ov["A17"] = "2. THỐNG KÊ & SO SÁNH PHIÊN BẢN CỬA HÀNG APP STORE"
        ws_ov["A17"].font = Font(bold=True, size=11, color="7D3C98")
        
        write_comparison_table(ws_ov, 18, "App Store", as_v2_stats, as_v3_stats, "7D3C98")
        
        ws_ov["A25"] = "PHÂN TÍCH & ĐÁNH GIÁ CHI TIẾT APP STORE"
        ws_ov["A25"].font = Font(bold=True, size=11, color="2C3E50")
        
        ws_ov.merge_cells("A26:F28")
        ws_ov["A26"] = as_text
        ws_ov["A26"].alignment = Alignment(wrap_text=True, vertical="top")
        ws_ov["A26"].font = Font(size=10)
        ws_ov["A26"].fill = PatternFill("solid", fgColor="F5EEF8")
        
        # Recommendations Section
        ws_ov["A30"] = "3. KHUYẾN NGHỊ VẬN HÀNH & TỐI ƯU TRẢI NGHIỆM PHIÊN BẢN MỚI (VERSION 3)"
        ws_ov["A30"].font = Font(bold=True, size=11, color="C0392B")
        
        ws_ov.merge_cells("A31:F36")
        ws_ov["A31"] = rec_text
        ws_ov["A31"].alignment = Alignment(wrap_text=True, vertical="top")
        ws_ov["A31"].font = Font(size=10, color="1C2833")
        ws_ov["A31"].fill = PatternFill("solid", fgColor="FDEDEC")
        
        # Overview layout structure
        ws_ov.column_dimensions["A"].width = 32
        ws_ov.column_dimensions["B"].width = 28
        ws_ov.column_dimensions["C"].width = 28
        ws_ov.column_dimensions["D"].width = 15
        ws_ov.column_dimensions["E"].width = 15
        ws_ov.column_dimensions["F"].width = 15
        
        for r in range(13, 16):
            ws_ov.row_dimensions[r].height = 22
        for r in range(26, 29):
            ws_ov.row_dimensions[r].height = 22
        for r in range(31, 37):
            ws_ov.row_dimensions[r].height = 22

        # Write data tabs
        col_order = ["author", "rating", "content", "at", "Version Segment", "Region", "Store", "sentiment", "category", "all_categories"]
        master_df[col_order].to_excel(writer, sheet_name="Raw_Data_VN", index=False)
        bad_df[col_order].to_excel(writer, sheet_name="Negative_Alerts", index=False)
        summary_pivot.to_excel(writer, sheet_name="Version_Summary", index=False)
        sentiment_dist.to_excel(writer, sheet_name="Sentiment_Summary", index=False)
        rating_dist.to_excel(writer, sheet_name="Rating_Distribution", index=False)
        category_dist.to_excel(writer, sheet_name="Top_Issues", index=False)

        if not top_keywords_df.empty:
            top_keywords_df.to_excel(writer, sheet_name="Top_Negative_Keywords", index=False)

        if "Sheet" in wb.sheetnames:
            wb.remove(wb["Sheet"])

        # Sheet styling
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            if sheet_name == "Overview_&_Assessment":
                continue
            elif sheet_name in ["Raw_Data_VN", "Negative_Alerts"]:
                sentiment_col = col_order.index("sentiment") + 1
                style_raw_sheet(ws, sentiment_col)
            else:
                _style_header(ws)
            _auto_column_width(ws)

    log.info(f"✅ Đã lưu file Excel: {output_path}")

def get_top_negative_keywords(master_df: pd.DataFrame, top_n: int = 15) -> pd.DataFrame:
    neg_df = master_df[master_df["sentiment"].isin(["Tiêu cực", "Hỗn hợp"])].copy()
    if neg_df.empty:
        return pd.DataFrame(columns=["keyword", "count", "lang", "example"])

    text_series = neg_df["content"].fillna("").str.lower()
    stats = []

    for lang_name, lang_dict in KEYWORD_DICT.items():
        for raw_pat in lang_dict["negative"]:
            pat = re.compile(raw_pat, re.IGNORECASE)
            matches = text_series.str.contains(pat, na=False)
            count = int(matches.sum())
            if count > 0:
                example_idx = matches[matches].index[0]
                example_text = neg_df.loc[example_idx, "content"]
                example = (example_text[:100] + "...") if len(example_text) > 100 else example_text
                display = re.sub(r'\\b|\\\(.*?\\\)|\^|\$', '', raw_pat).strip()
                stats.append({
                    "keyword": display or raw_pat,
                    "count": count,
                    "lang": lang_name.upper(),
                    "example": example
                })

    df_kw = pd.DataFrame(stats)
    if not df_kw.empty:
        df_kw = df_kw.sort_values("count", ascending=False).head(top_n).reset_index(drop=True)
    return df_kw

# ─────────────────────────────────────────────────────────────────
# 5. MAIN EXECUTION
# ─────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    log.info("═" * 60)
    log.info("  Crossfire Legends Review Scraper via Sensor Tower (VN ONLY) STARTED")
    log.info("═" * 60)
    
    if not SENSORTOWER_API_KEY or SENSORTOWER_API_KEY == "YOUR_SENSORTOWER_API_KEY_HERE":
        log.error("SENSORTOWER_API_KEY is not defined in key_api.env! Exiting.")
        sys.exit(1)
        
    all_dfs = []
    
    # Loop through VN region configuration
    for region in GOOGLE_APPS.keys():
        gp_id = GOOGLE_APPS[region]
        apple_id = APPLE_APPS[region]
        
        # 1. Fetch Google Play reviews
        df_gp = fetch_sensortower_reviews("android", gp_id, region)
        if not df_gp.empty:
            all_dfs.append(df_gp)
            
        # 2. Fetch App Store reviews
        df_ios = fetch_sensortower_reviews("ios", apple_id, region)
        if not df_ios.empty:
            all_dfs.append(df_ios)
            
    if not all_dfs:
        log.error("❌ Không lấy được dữ liệu nào từ Sensor Tower API.")
    else:
        master_df = pd.concat(all_dfs, ignore_index=True)
        log.info(f"Tổng raw: {len(master_df)} reviews trước dedup.")

        # Dedup
        master_df.drop_duplicates(subset=["author", "content", "at"], inplace=True)
        log.info(f"Sau dedup: {len(master_df)} reviews.")

        # Sanitize columns starting with '=' to prevent Excel formula corruption
        for col in ["author", "content"]:
            if col in master_df.columns:
                master_df[col] = master_df[col].fillna("").astype(str).apply(
                    lambda x: f" {x}" if x.startswith("=") else x
                )

        master_df["at"] = (
            pd.to_datetime(master_df["at"], errors="coerce")
              .dt.tz_localize(None)
        )

        # Segment reviews into Version 2 & Version 3
        master_df["Version Segment"] = master_df["at"].apply(assign_version_segment)
        
        # Keep only reviews from March 3rd onwards
        master_df = master_df[master_df["Version Segment"].isin([
            "Version 2 (03/03 - 27/04)",
            "Version 3 (28/04 - Present)"
        ])].copy()
        
        log.info(f"Sau khi lọc ngày (từ 03/03): {len(master_df)} reviews.")

        # Apply analysis
        master_df["sentiment"] = master_df.apply(analyze_sentiment, axis=1)
        master_df = categorize_feedback(master_df)

        top_keywords_df = get_top_negative_keywords(master_df)

        bad_df = master_df[
            (master_df["rating"] <= 2) | (master_df["sentiment"] == "Tiêu cực")
        ].copy()

        log.info(f"Negative/Low-rating reviews: {len(bad_df)}")

        # Write report
        write_excel(master_df, bad_df, top_keywords_df, OUTPUT_FILE)

        print("\n" + "═" * 50)
        print(f"  TỔNG REVIEWS (SENSOR TOWER VN): {len(master_df)}")
        print(f"  CẦN CHÚ Ý                     : {len(bad_df)}")
        print(f"  TOP KEYWORDS                  : {len(top_keywords_df)} từ tiêu cực")
        print("═" * 50)
        print(f"  File: {OUTPUT_FILE}")
        print("═" * 50 + "\n")
